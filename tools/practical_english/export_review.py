#!/usr/bin/env python3
"""把句庫草稿匯出成外部審閱用的 Excel 與 CSV（不動 pe_core.json）。

來源：source/draft_*.tsv、sense_review.tsv、assets/data 的 4 份詞表。
輸出：<輸出資料夾>/pe_core_review.xlsx（說明／例句／詞義連結 3 個工作表）
      <輸出資料夾>/pe_core_review.csv（例句工作表，UTF-8 BOM，Excel 可直接開）

用法：python3 tools/practical_english/export_review.py <輸出資料夾>
"""
import csv, glob, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
FILES = ['ngsl_2809', 'ngsl_spoken_720', 'phrase_list_506', 'phave_list_150']
LIST_NAME = {'ngsl_2809': 'NGSL', 'ngsl_spoken_720': 'NGSL 口語', 'phrase_list_506': 'PHRASE List',
             'phave_list_150': 'PhaVE'}
REVIEW_COLS = ['審閱結果', '建議英文', '建議中文', '審閱意見']
MAIN_COLS = ['句子編號', '主要詞', '主要詞 ID', '詞表', '詞表中文', '英文例句', '繁中翻譯', '難度', '類別',
             '同句連結的同拼字詞 ID', '次要詞', '次要詞 ID', '片語', '備註'] + REVIEW_COLS


def read_tsv(path):
    with open(path, encoding='utf-8') as f:
        return list(csv.DictReader(f, delimiter='\t', quoting=csv.QUOTE_NONE))


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    out_dir = sys.argv[1]
    os.makedirs(out_dir, exist_ok=True)
    items = {}
    for name in FILES:
        for it in json.load(open(os.path.join(ROOT, 'assets', 'data', name + '.json'), encoding='utf-8'))['items']:
            items[it['id']] = it
    sense = read_tsv(os.path.join(HERE, 'sense_review.tsv'))
    linked = {}
    for r in sense:
        if r['decision'] == 'include':
            linked.setdefault(r['sentence_id'], []).append(r['candidate_id'])

    rows, sentences = [], {}
    for path in sorted(glob.glob(os.path.join(HERE, 'source', 'draft_*.tsv'))):
        for r in read_tsv(path):
            pid = r['primary_word_id']
            same = [c for c in items if c != pid and items[c]['w'].strip().lower() == r['word'].strip().lower()
                    and items[c]['m'].get('zh-TW') == items[pid]['m'].get('zh-TW')]
            secs = [s for s in r['secondary_word_ids'].split('|') if s]
            sentences[r['id']] = r
            rows.append([
                f"pe_core_{int(r['id']):06d}", r['word'], pid, LIST_NAME[pid.rsplit('_', 1)[0]],
                items[pid]['m'].get('zh-TW', ''), r['sentence'], r['zh-TW'], r['level'], r['category'],
                ' '.join(same + linked.get(r['id'], [])),
                ', '.join(items[s]['w'] for s in secs), ' '.join(secs),
                r['phrase_candidate'], r['tagging_notes'], '', '', '', ''])

    wb = Workbook()
    info = wb.active
    info.title = '說明'
    for line in [
        ['智慧聽覺巡航 V2 實用英文：內建例句審閱表（草稿，尚未審核）'],
        [f'例句 {len(rows)} 句，涵蓋 4 份詞表共 {len(items)} 個詞（NGSL 2,809、NGSL 口語 720、PHRASE List 506、PhaVE 150）。'],
        ['每個詞至少出現在一句的「主要詞」或「同句連結的同拼字詞 ID」裡。拼字相同、詞義也相同的詞共用一句；詞義不同的各自一句。'],
        [''],
        ['請審閱的重點'],
        ['1. 英文是否自然、正確，母語人士會不會這樣說。'],
        ['2. 繁中翻譯是否正確、符合台灣用語。'],
        ['3. 主要詞在句中的意思，是否符合「詞表中文」。備註欄有寫「在此為…」的，表示句中用法與詞表中文不完全相同，請特別確認。'],
        ['4. 內容是否適合所有年齡（不雅、敏感、爭議）。'],
        ['5. 難度（A1–C1）與類別是否合理。'],
        [''],
        ['填寫方式'],
        ['「審閱結果」選 OK／修改／刪除。選「修改」時，在「建議英文」「建議中文」填修改後的句子；其他意見寫在「審閱意見」。'],
        ['請不要更改「句子編號」「主要詞 ID」等前面的欄位，也不要調整列的順序。'],
        [''],
        ['「詞義連結」工作表'],
        ['同一個拼字在兩份詞表都有、但中文解釋不同時，判斷這句是否也算那個詞的例句。include＝算；exclude＝不算（該詞另有自己的例句）。'],
        ['如不同意建議，請在「審閱結果」填 include 或 exclude 並寫原因。'],
    ]:
        info.append(line)
    info['A1'].font = Font(bold=True, size=14)
    for c in ('A5', 'A12', 'A16'):
        info[c].font = Font(bold=True)
    info.column_dimensions['A'].width = 120

    head_fill = PatternFill('solid', fgColor='5B8A72')
    review_fill = PatternFill('solid', fgColor='FFF7E0')

    def sheet(title, cols, data, widths, review_from):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for r in data:
            ws.append(r)
        for i, c in enumerate(ws[1], 1):
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
            for c in row[review_from:]:
                c.fill = review_fill
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        return ws

    ws = sheet('例句', MAIN_COLS, rows,
               [15, 14, 22, 11, 18, 44, 30, 6, 10, 22, 24, 30, 16, 28, 10, 40, 30, 30], len(MAIN_COLS) - 4)
    dv = DataValidation(type='list', formula1='"OK,修改,刪除"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'O2:O{len(rows) + 1}')

    sense_rows = []
    for r in sense:
        s = sentences.get(r['sentence_id'])
        if not s:
            continue
        sense_rows.append([f"pe_core_{int(r['sentence_id']):06d}", s['word'], s['sentence'], s['zh-TW'],
                           r['primary_word_id'], r['primary_zh'], r['candidate_id'], r['candidate_zh'],
                           r['decision'], r['note'], '', ''])
    ws2 = sheet('詞義連結', ['句子編號', '單字', '英文例句', '繁中翻譯', '主要詞 ID', '主要詞中文', '候選 ID',
                          '候選中文', '建議', '理由', '審閱結果', '審閱意見'],
                sense_rows, [15, 12, 40, 28, 22, 16, 22, 22, 9, 36, 10, 30], 10)
    dv2 = DataValidation(type='list', formula1='"include,exclude"', allow_blank=True)
    ws2.add_data_validation(dv2)
    dv2.add(f'K2:K{len(sense_rows) + 1}')

    xlsx = os.path.join(out_dir, 'pe_core_review.xlsx')
    wb.save(xlsx)
    with open(os.path.join(out_dir, 'pe_core_review.csv'), 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(MAIN_COLS)
        w.writerows(rows)
    print(f'例句 {len(rows)} 句、詞義連結 {len(sense_rows)} 列 → {xlsx}')


if __name__ == '__main__':
    main()
